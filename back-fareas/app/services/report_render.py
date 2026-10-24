import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape, portrait
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

INSTITUCION = "FACULTAD DE INFORMÁTICA - UNAMBA"
SISTEMA = "Sistema Fareas · Asistencia con reconocimiento facial"
AZUL = "155DFC"

TONOS = {
    "ok": "15803D",
    "warn": "B45309",
    "bad": "B91C1C",
    "info": "0369A1",
    "muted": "6B7280",
}
FONDOS = {
    "ok": "DCFCE7",
    "warn": "FEF3C7",
    "bad": "FEE2E2",
    "info": "E0F2FE",
    "muted": "F3F4F6",
}

FORMATOS = {
    "pdf": ("application/pdf", "pdf"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "csv": ("text/csv; charset=utf-8", "csv"),
}


@dataclass(frozen=True)
class Celda:
    valor: object
    tono: str | None = None


@dataclass(frozen=True)
class Columna:
    clave: str
    titulo: str
    ancho: float = 1.0
    alinear: str = "left"
    numerica: bool = False


@dataclass(frozen=True)
class Indicador:
    etiqueta: str
    valor: str
    tono: str = "muted"


@dataclass
class Tabla:
    clave: str
    titulo: str
    descripcion: str
    columnas: list[Columna]
    filas: list[list[Celda]]
    filtros: list[tuple[str, str]] = field(default_factory=list)
    indicadores: list[Indicador] = field(default_factory=list)
    horizontal: bool = True
    truncada: bool = False

    def valores(self) -> list[list[object]]:
        return [[celda.valor for celda in fila] for fila in self.filas]


def texto(valor: object) -> str:
    if valor is None:
        return ""
    return str(valor)


def nombre_archivo(tabla: Tabla, formato: str, ahora: datetime | None = None) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", tabla.clave.lower()).strip("-")
    return f"{base}_{(ahora or datetime.now()):%Y%m%d_%H%M}.{FORMATOS[formato][1]}"


def a_csv(tabla: Tabla) -> bytes:
    salida = io.StringIO()
    escritor = csv.writer(salida)
    escritor.writerow([columna.titulo for columna in tabla.columnas])
    for fila in tabla.valores():
        escritor.writerow([texto(valor) for valor in fila])
    return b"\xef\xbb\xbf" + salida.getvalue().encode("utf-8")


def a_xlsx(tabla: Tabla) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = re.sub(r"[\[\]:*?/\\]", "", tabla.titulo)[:31] or "Reporte"
    total = max(len(tabla.columnas), 4)
    fino = Side(style="thin", color="D1D5DB")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)

    def combinada(fila: int, contenido: str, fuente: Font, alto: float | None = None) -> None:
        hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=total)
        celda = hoja.cell(fila, 1, contenido)
        celda.font = fuente
        celda.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        if alto:
            hoja.row_dimensions[fila].height = alto

    combinada(1, INSTITUCION, Font(bold=True, size=13, color=AZUL), 22)
    combinada(2, tabla.titulo, Font(bold=True, size=15, color="111827"), 24)
    combinada(3, tabla.descripcion, Font(size=10, color="6B7280"))
    fila = 4
    if tabla.filtros:
        combinada(
            fila,
            "Filtros: " + "  ·  ".join(f"{k}: {v}" for k, v in tabla.filtros),
            Font(size=10, color="374151"),
        )
        fila += 1
    combinada(
        fila,
        f"Generado: {datetime.now():%d/%m/%Y %H:%M}  ·  Registros: {len(tabla.filas)}"
        + ("  ·  Resultado limitado a los primeros registros" if tabla.truncada else ""),
        Font(size=9, color="9CA3AF"),
    )
    fila += 2

    if tabla.indicadores:
        for indice, indicador in enumerate(tabla.indicadores):
            columna = indice + 1
            etiqueta = hoja.cell(fila, columna, indicador.etiqueta)
            etiqueta.font = Font(size=9, color="6B7280")
            etiqueta.alignment = Alignment(horizontal="center")
            valor = hoja.cell(fila + 1, columna, indicador.valor)
            valor.font = Font(bold=True, size=14, color=TONOS[indicador.tono])
            valor.alignment = Alignment(horizontal="center")
            for destino in (etiqueta, valor):
                destino.fill = PatternFill("solid", fgColor=FONDOS[indicador.tono])
        fila += 3

    cabecera = fila
    for indice, columna in enumerate(tabla.columnas, start=1):
        celda = hoja.cell(cabecera, indice, columna.titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = borde
    hoja.row_dimensions[cabecera].height = 24

    for desplazamiento, datos in enumerate(tabla.filas, start=1):
        for indice, (columna, dato) in enumerate(zip(tabla.columnas, datos), start=1):
            celda = hoja.cell(cabecera + desplazamiento, indice, dato.valor)
            celda.border = borde
            celda.alignment = Alignment(horizontal=columna.alinear, vertical="center", wrap_text=True)
            if dato.tono:
                celda.font = Font(bold=True, color=TONOS[dato.tono])
                celda.fill = PatternFill("solid", fgColor=FONDOS[dato.tono])
            if isinstance(dato.valor, float):
                celda.number_format = "0.0"

    ancho_total = sum(columna.ancho for columna in tabla.columnas) or 1
    for indice, columna in enumerate(tabla.columnas, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = max(
            9, round(columna.ancho / ancho_total * 140)
        )
    hoja.freeze_panes = hoja.cell(cabecera + 1, 1)
    if tabla.filas:
        hoja.auto_filter.ref = f"A{cabecera}:{get_column_letter(len(tabla.columnas))}{cabecera + len(tabla.filas)}"
    hoja.sheet_view.showGridLines = False
    hoja.page_setup.orientation = "landscape" if tabla.horizontal else "portrait"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.print_title_rows = f"{cabecera}:{cabecera}"

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


class _NumeroDePagina(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._paginas: list[dict] = []

    def showPage(self):
        self._paginas.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._paginas)
        for estado in self._paginas:
            self.__dict__.update(estado)
            self.setFont("Helvetica", 7.5)
            self.setFillColor(colors.HexColor("#9CA3AF"))
            ancho = self._pagesize[0]
            self.drawString(12 * mm, 8 * mm, f"{SISTEMA}")
            self.drawRightString(ancho - 12 * mm, 8 * mm, f"Página {self._pageNumber} de {total}")
            super().showPage()
        super().save()


def a_pdf(tabla: Tabla) -> bytes:
    tamano = landscape(A4) if tabla.horizontal else portrait(A4)
    salida = io.BytesIO()
    documento = SimpleDocTemplate(
        salida, pagesize=tamano, topMargin=12 * mm, bottomMargin=14 * mm,
        leftMargin=12 * mm, rightMargin=12 * mm, title=f"{INSTITUCION} - {tabla.titulo}",
        author="Sistema Fareas",
    )
    estilos = getSampleStyleSheet()
    institucion = ParagraphStyle(
        "i", parent=estilos["Normal"], fontSize=9, textColor=colors.HexColor(f"#{AZUL}"),
        fontName="Helvetica-Bold",
    )
    titulo = ParagraphStyle(
        "t", parent=estilos["Title"], fontSize=17, alignment=0, spaceAfter=1, leading=20,
        textColor=colors.HexColor("#111827"),
    )
    suave = ParagraphStyle(
        "s", parent=estilos["Normal"], fontSize=8.5, textColor=colors.HexColor("#6B7280"), leading=11
    )
    celda = ParagraphStyle("c", parent=estilos["Normal"], fontSize=7.5, leading=9)
    valor_estilo = ParagraphStyle("v", parent=celda, leading=20)
    cabecera = ParagraphStyle(
        "h", parent=celda, textColor=colors.white, fontName="Helvetica-Bold"
    )

    historia: list = [
        Paragraph(INSTITUCION, institucion),
        Paragraph(tabla.titulo, titulo),
        Paragraph(tabla.descripcion, suave),
    ]
    if tabla.filtros:
        historia.append(
            Paragraph(
                "<b>Filtros:</b> " + "  ·  ".join(f"{k}: {v}" for k, v in tabla.filtros), suave
            )
        )
    historia.append(
        Paragraph(
            f"Generado el {datetime.now():%d/%m/%Y a las %H:%M}  ·  Registros: {len(tabla.filas)}"
            + ("  ·  Resultado limitado a los primeros registros" if tabla.truncada else ""),
            suave,
        )
    )
    historia.append(Spacer(1, 4 * mm))

    ancho_util = tamano[0] - 24 * mm
    if tabla.indicadores:
        ancho_tarjeta = ancho_util / len(tabla.indicadores)
        etiquetas = [
            Paragraph(f'<font size="7" color="#6B7280">{i.etiqueta}</font>', celda)
            for i in tabla.indicadores
        ]
        valores = [
            Paragraph(f'<font size="15" color="#{TONOS[i.tono]}"><b>{i.valor}</b></font>', valor_estilo)
            for i in tabla.indicadores
        ]
        tarjetas = Table([etiquetas, valores], colWidths=[ancho_tarjeta] * len(tabla.indicadores))
        estilo_tarjetas = [
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.white),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
        for indice, indicador in enumerate(tabla.indicadores):
            estilo_tarjetas.append(
                ("BACKGROUND", (indice, 0), (indice, 1), colors.HexColor(f"#{FONDOS[indicador.tono]}"))
            )
        tarjetas.setStyle(TableStyle(estilo_tarjetas))
        historia += [tarjetas, Spacer(1, 4 * mm)]

    pesos = [columna.ancho for columna in tabla.columnas]
    anchos = [ancho_util * peso / sum(pesos) for peso in pesos]
    alineacion = {"left": 0, "center": 1, "right": 2}
    datos = [[Paragraph(columna.titulo, cabecera) for columna in tabla.columnas]]
    for fila in tabla.filas:
        renglon = []
        for columna, dato in zip(tabla.columnas, fila):
            estilo = ParagraphStyle(
                f"c{columna.clave}", parent=celda, alignment=alineacion.get(columna.alinear, 0)
            )
            contenido = texto(dato.valor).replace("&", "&amp;").replace("<", "&lt;")
            if dato.tono:
                contenido = f'<font color="#{TONOS[dato.tono]}"><b>{contenido}</b></font>'
            renglon.append(Paragraph(contenido, estilo))
        datos.append(renglon)
    if not tabla.filas:
        datos.append([Paragraph("Sin registros para los filtros elegidos.", celda)] + [""] * (len(tabla.columnas) - 1))

    cuerpo = Table(datos, colWidths=anchos, repeatRows=1)
    estilo_cuerpo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{AZUL}")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]
    if not tabla.filas:
        estilo_cuerpo.append(("SPAN", (0, 1), (-1, 1)))
    cuerpo.setStyle(TableStyle(estilo_cuerpo))
    historia.append(cuerpo)
    documento.build(historia, canvasmaker=_NumeroDePagina)
    return salida.getvalue()


def renderizar(tabla: Tabla, formato: str) -> bytes:
    if formato == "xlsx":
        return a_xlsx(tabla)
    if formato == "pdf":
        return a_pdf(tabla)
    return a_csv(tabla)
