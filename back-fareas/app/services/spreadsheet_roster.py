from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass

from openpyxl import load_workbook

from app.services.ocr_roster import (
    DOMINIO_POR_DEFECTO,
    MAX_FILAS,
    AlumnoExtraido,
    _normalizar_ciclo,
    _tipo_columna,
    validar_alumno,
)

SPREADSHEET_EXTENSIONS = (".xlsx", ".xlsm", ".csv", ".tsv", ".txt", ".xls")
ZIP_MAGIC = b"PK\x03\x04"
OLE_MAGIC = b"\xd0\xcf\x11\xe0"
MAX_SCAN_ROWS = 5000
MAX_SCAN_COLUMNS = 60
HEADER_LOOKBACK = 3
MAX_NAME_COLUMNS = 3
NAME_TYPES = ("full_name", "apellidos", "ap_pat", "ap_mat", "nombres")
RE_CODE = re.compile(r"\d{6}")
RE_DNI = re.compile(r"\d{8}")
RE_MOBILE = re.compile(r"9\d{8}")
RE_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
RE_NAME = re.compile(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ][A-Za-zÁÉÍÓÚÜÑáéíóúüñ\s,.'’-]*")


@dataclass(slots=True)
class SheetRead:
    kind: str
    sheets: int
    students: list[AlumnoExtraido]


def is_spreadsheet(raw: bytes, filename: str | None) -> bool:
    name = (filename or "").lower()
    if name.endswith(SPREADSHEET_EXTENSIONS):
        return True
    return raw.startswith((ZIP_MAGIC, OLE_MAGIC))


def _cell_text(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, str):
        return re.sub(r"\s+", " ", value.replace("\xa0", " ")).strip()
    return ""


def _trim(rows: list[list[str]]) -> list[list[str]]:
    return [row[:MAX_SCAN_COLUMNS] for row in rows[:MAX_SCAN_ROWS]]


def _workbook_sheets(raw: bytes) -> list[list[list[str]]]:
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("No se pudo leer el Excel; ¿es un archivo .xlsx válido?") from exc
    sheets: list[list[list[str]]] = []
    try:
        for sheet in workbook.worksheets:
            rows = [
                [_cell_text(value) for value in row]
                for row in sheet.iter_rows(max_row=MAX_SCAN_ROWS, values_only=True)
            ]
            rows = _trim(rows)
            if any(any(cell for cell in row) for row in rows):
                sheets.append(rows)
    finally:
        workbook.close()
    return sheets


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _delimiter(text: str) -> str:
    sample = "\n".join(text.splitlines()[:20])
    counts = {candidate: sample.count(candidate) for candidate in (",", ";", "\t", "|")}
    best = max(counts, key=counts.get)
    return best if counts[best] else ","


def _csv_sheets(raw: bytes) -> list[list[list[str]]]:
    text = _decode(raw)
    reader = csv.reader(io.StringIO(text), delimiter=_delimiter(text))
    rows = _trim([[_cell_text(cell) for cell in row] for row in reader])
    return [rows] if any(any(cell for cell in row) for row in rows) else []


def _is_key(cell: str) -> bool:
    return bool(RE_EMAIL.search(cell) or RE_CODE.fullmatch(cell) or RE_DNI.fullmatch(cell))


def _clean_mobile(cell: str) -> str:
    compact = re.sub(r"[\s.()-]", "", cell)
    return compact[3:] if compact.startswith("+51") else compact


def _fraction(cells: list[str], predicate) -> float:
    return sum(1 for cell in cells if predicate(cell)) / len(cells) if cells else 0.0


def _infer_type(cells: list[str]) -> str | None:
    cells = [cell for cell in cells if cell]
    if not cells:
        return None
    if _fraction(cells, lambda c: bool(RE_EMAIL.search(c))) >= 0.5:
        return "email"
    if _fraction(cells, lambda c: bool(RE_CODE.fullmatch(c))) >= 0.5:
        return "code"
    if _fraction(cells, lambda c: bool(RE_DNI.fullmatch(c))) >= 0.5:
        return "dni"
    if _fraction(cells, lambda c: bool(RE_MOBILE.fullmatch(_clean_mobile(c)))) >= 0.5:
        return "whatsapp"
    if _fraction(cells, lambda c: len(c) <= 10 and _normalizar_ciclo(c) is not None) >= 0.5:
        return "semester"
    if _fraction(cells, lambda c: bool(RE_NAME.fullmatch(c))) >= 0.6:
        if len(cells) > 3 and len(set(cells)) <= 1:
            return None
        return "text"
    return None


def _header_types(row: list[str]) -> dict[int, str]:
    types: dict[int, str] = {}
    seen: set[str] = set()
    for index, cell in enumerate(row):
        if not cell or _is_key(cell):
            continue
        kind = _tipo_columna(cell)
        if kind is not None and kind not in seen:
            seen.add(kind)
            types[index] = kind
    return types


def _find_header(rows: list[list[str]], first_data: int) -> dict[int, str]:
    for offset in range(1, HEADER_LOOKBACK + 1):
        index = first_data - offset
        if index < 0:
            break
        row = rows[index]
        if any(_is_key(cell) for cell in row):
            break
        types = _header_types(row)
        if len(types) >= 2 or (types and offset == 1):
            return types
    return {}


def _column_types(
    rows: list[list[str]], data: list[int], header: dict[int, str]
) -> dict[int, str]:
    width = max(len(rows[i]) for i in data)
    types = dict(header)
    inferred_text: list[int] = []
    for column in range(width):
        if column in types:
            continue
        values = [rows[i][column] if column < len(rows[i]) else "" for i in data]
        kind = _infer_type(values)
        if kind == "text":
            inferred_text.append(column)
        elif kind is not None and kind not in types.values():
            types[column] = kind
    if not any(kind in NAME_TYPES for kind in types.values()) and inferred_text:
        block = [inferred_text[0]]
        for column in inferred_text[1:]:
            if column == block[-1] + 1 and len(block) < MAX_NAME_COLUMNS:
                block.append(column)
            else:
                break
        if len(block) == 1:
            types[block[0]] = "full_name"
        else:
            for column in block[:-1]:
                types[column] = "apellidos"
            types[block[-1]] = "nombres"
    return types


def _clean_name(text: str) -> str:
    text = re.sub(r"[^A-Za-zÁÉÍÓÚÜÑáéíóúüñ\s,.'’-]", " ", text)
    text = re.sub(r"\s+", " ", text).strip(" -–-:;,.")
    return re.sub(r"\s+,", ",", text)


def _full_name(row: list[str], types: dict[int, str]) -> str:
    parts: dict[str, list[str]] = {kind: [] for kind in NAME_TYPES}
    for column in sorted(types):
        kind = types[column]
        if kind in parts and column < len(row) and row[column]:
            parts[kind].append(row[column])
    if parts["full_name"]:
        return _clean_name(" ".join(parts["full_name"]))
    surnames = " ".join(parts["ap_pat"] + parts["ap_mat"] + parts["apellidos"])
    given = " ".join(parts["nombres"])
    surnames, given = _clean_name(surnames), _clean_name(given)
    if surnames and given:
        return f"{surnames}, {given}"
    return surnames or given


def _value(row: list[str], types: dict[int, str], kind: str) -> str:
    for column, current in types.items():
        if current == kind and column < len(row):
            return row[column]
    return ""


def _dni_digits(value: str) -> str | None:
    compact = re.sub(r"\D", "", value)
    if not compact:
        return None
    return compact.zfill(8) if len(compact) == 7 else compact


def _student_from_row(
    row: list[str], types: dict[int, str], number: int, sheet: int
) -> AlumnoExtraido | None:
    email_match = RE_EMAIL.search(_value(row, types, "email")) or next(
        (m for cell in row if (m := RE_EMAIL.search(cell))), None
    )
    email = email_match.group(0) if email_match else None
    code = _value(row, types, "code")
    code = re.sub(r"\D", "", code) if code else ""
    if not code and email:
        local = email.split("@", 1)[0]
        code = local if RE_CODE.fullmatch(local) else ""
    dni_raw = _value(row, types, "dni")
    dni = _dni_digits(dni_raw)
    if not (code or dni or email):
        return None
    mobile_raw = _value(row, types, "whatsapp")
    mobile = _clean_mobile(mobile_raw) if mobile_raw else ""
    cycle_raw = _value(row, types, "semester")
    return AlumnoExtraido(
        fila=number,
        code=code or None,
        dni=dni,
        full_name=_full_name(row, types) or None,
        email=email,
        semester=_normalizar_ciclo(cycle_raw) if cycle_raw else None,
        whatsapp=mobile if RE_MOBILE.fullmatch(mobile) else None,
        linea=" | ".join(cell for cell in row if cell)[:200],
        pagina=sheet,
    )


def _parse_sheet(
    rows: list[list[str]],
    *,
    default_semester: int,
    domain: str,
    sheet: int,
    seen: tuple[set[str], set[str], set[str]],
    start: int,
) -> list[AlumnoExtraido]:
    data = [index for index, row in enumerate(rows) if any(_is_key(cell) for cell in row)]
    if not data:
        return []
    types = _column_types(rows, data, _find_header(rows, data[0]))
    students: list[AlumnoExtraido] = []
    for index in data:
        student = _student_from_row(rows[index], types, start + len(students) + 1, sheet)
        if student is None:
            continue
        if student.semester is None:
            student.nota = "Ciclo no detectado; se usó el ciclo por defecto"
        validar_alumno(
            student,
            default_semester=default_semester,
            dominio=domain,
            vistos_codigo=seen[0],
            vistos_dni=seen[1],
            vistos_email=seen[2],
        )
        students.append(student)
    return students


def read_spreadsheet(
    raw: bytes,
    filename: str | None,
    *,
    default_semester: int,
    domain: str = DOMINIO_POR_DEFECTO,
) -> SheetRead:
    if not raw:
        raise ValueError("El archivo está vacío")
    name = (filename or "").lower()
    if raw.startswith(OLE_MAGIC) or name.endswith(".xls"):
        raise ValueError(
            "Formato no soportado: guarda el .xls como .xlsx o CSV desde Excel e inténtalo de nuevo"
        )
    if raw.startswith(ZIP_MAGIC):
        kind, sheets = "xlsx", _workbook_sheets(raw)
    else:
        kind, sheets = "csv", _csv_sheets(raw)

    seen: tuple[set[str], set[str], set[str]] = (set(), set(), set())
    students: list[AlumnoExtraido] = []
    for number, rows in enumerate(sheets, start=1):
        students.extend(
            _parse_sheet(
                rows,
                default_semester=default_semester,
                domain=domain,
                sheet=number,
                seen=seen,
                start=len(students),
            )
        )
        if len(students) > MAX_FILAS:
            raise ValueError(f"La nómina supera el máximo de {MAX_FILAS} filas")
    return SheetRead(kind=kind, sheets=len(sheets), students=students)
