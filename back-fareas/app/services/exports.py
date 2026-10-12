"""Reportes de asistencia — RF-16: Excel (openpyxl) y PDF (reportlab).

Funciones PURAS: reciben los datos ya consultados y devuelven los bytes del
archivo; los routers solo hacen la consulta y arman la respuesta HTTP con el
Content-Disposition. Sin dependencia del request ni de la sesión de BD.
"""
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

# Encabezado institucional compartido por ambos formatos (RF-16)
TITULO = "FACULTAD DE INFORMÁTICA - UNAMBA"
SUBTITULO = "Reporte de Asistencia - Sistema Fareas"

AZUL = "155DFC"          # --fareas-primary
AZUL_CLARO = "E8F0FE"

# --- Columnas estándar (misma tabla en pantalla, Excel y PDF) ----------------
HEADERS = ["Estudiante", "Código", "Curso", "Grupo", "Aula", "Estado", "Fecha", "Hora", "Método"]
ANCHOS_PDF = [58 * mm, 18 * mm, 58 * mm, 12 * mm, 20 * mm, 22 * mm, 24 * mm, 16 * mm, 18 * mm]

COLORES_ESTADO = {
    "asistio": ("22C55E", "Asistió"),
    "tardanza": ("F59E0B", "Tardanza"),
    "falta": ("EF4444", "Falta"),
    "falta_justificada": ("0EA5E9", "Falta justificada"),  # R6
    "sin_registro": ("9CA3AF", "Sin registro"),
}


def _estado_de(r: dict) -> str:
    """Clave de color/etiqueta: R6 distingue la falta JUSTIFICADA (no cuenta DPI)."""
    estado = str(r.get("status", ""))
    if estado == "falta" and r.get("is_justified"):
        return "falta_justificada"
    return estado


def _fila(r: dict) -> list[str]:
    hora = (r.get("marked_at") or "")[11:16] if r.get("marked_at") else "-"
    return [
        str(r.get("student_name", "")),
        str(r.get("student_code") or "-"),
        str(r.get("course_name", "")),
        str(r.get("group_code", "-")),
        str(r.get("room_code", "")),
        COLORES_ESTADO.get(_estado_de(r), ("", str(r.get("status", "-"))))[1],
        str(r.get("att_date") or ""),
        hora,
        str(r.get("method", "")),
    ]


def _nombre_archivo(ext: str, desde: date | None, hasta: date | None, now: datetime) -> str:
    rango = f"_{desde.isoformat()}_{hasta.isoformat()}" if desde or hasta else ""
    return f"asistencia{rango}_{now:%Y%m%d_%H%M}.{ext}"


def nombre_excel(desde: date | None, hasta: date | None, now: datetime | None = None) -> str:
    return _nombre_archivo("xlsx", desde, hasta, now or datetime.now())


def nombre_pdf(desde: date | None, hasta: date | None, now: datetime | None = None) -> str:
    return _nombre_archivo("pdf", desde, hasta, now or datetime.now())


# ---------------------------------------------------------------------------
# Excel (openpyxl)
# ---------------------------------------------------------------------------
def excel_asistencia(filas: list[dict], titulo_periodo: str | None = None) -> bytes:
    """Libro de una hoja: encabezado institucional + tabla con estilos."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencia"

    # Encabezado institucional
    ws.merge_cells("A1:I1")
    c = ws["A1"]
    c.value = TITULO
    c.font = Font(bold=True, size=13, color=AZUL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A2:I2")
    c = ws["A2"]
    c.value = f"{SUBTITULO}{f' - {titulo_periodo}' if titulo_periodo else ''}"
    c.font = Font(bold=True, size=11, color="374151")
    c.alignment = Alignment(horizontal="center")
    ws.merge_cells("A3:I3")
    ws["A3"] = f"Generado: {datetime.now():%d/%m/%Y %H:%M}  ·  Registros: {len(filas)}"
    ws["A3"].font = Font(size=9, color="9CA3AF")
    ws["A3"].alignment = Alignment(horizontal="center")

    # Fila de encabezados
    fila_hdr = 5
    fill_hdr = PatternFill("solid", fgColor=AZUL)
    fino = Side(style="thin", color="D1D5DB")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    for col, titulo in enumerate(HEADERS, start=1):
        c = ws.cell(row=fila_hdr, column=col, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill_hdr
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = borde

    # Datos con color por estado
    fills_estado = {k: PatternFill("solid", fgColor=f"{v[0]}22") for k, v in COLORES_ESTADO.items()}
    for i, r in enumerate(filas):
        valores = _fila(r)
        for col, valor in enumerate(valores, start=1):
            c = ws.cell(row=fila_hdr + 1 + i, column=col, value=valor)
            c.border = borde
            if col == 6:  # Estado
                estado = _estado_de(r)
                c.font = Font(bold=True, color=COLORES_ESTADO.get(estado, ("111827",))[0])
                fill = fills_estado.get(estado)
                if fill:
                    c.fill = fill

    # Anchos y freeze
    anchos = [32, 10, 38, 8, 12, 12, 12, 8, 10]
    for col, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.freeze_panes = f"A{fila_hdr + 1}"
    ws.auto_filter.ref = f"A{fila_hdr}:I{fila_hdr + len(filas)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF (reportlab)
# ---------------------------------------------------------------------------
def pdf_asistencia(filas: list[dict], titulo_periodo: str | None = None) -> bytes:
    """PDF horizontal con encabezado institucional y tabla paginada."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        topMargin=14 * mm, bottomMargin=12 * mm, leftMargin=12 * mm, rightMargin=12 * mm,
        title=f"{TITULO} - {SUBTITULO}", author="Sistema Fareas",
    )
    styles = getSampleStyleSheet()
    st_titulo = ParagraphStyle("titulo", parent=styles["Title"], fontSize=13,
                               textColor=colors.HexColor(f"#{AZUL}"), spaceAfter=1 * mm)
    st_sub = ParagraphStyle("sub", parent=styles["Normal"], fontSize=10, alignment=1,
                            textColor=colors.HexColor("#374151"))
    st_pie = ParagraphStyle("pie", parent=styles["Normal"], fontSize=7.5,
                            alignment=1, textColor=colors.HexColor("#9CA3AF"))
    st_celda = ParagraphStyle("celda", parent=styles["Normal"], fontSize=7, leading=8.5)

    story: list = [
        Paragraph(TITULO, st_titulo),
        Paragraph(f"{SUBTITULO}{f' - {titulo_periodo}' if titulo_periodo else ''}", st_sub),
        Spacer(1, 1.5 * mm),
        Paragraph(f"Generado: {datetime.now():%d/%m/%Y %H:%M} · Registros: {len(filas)}", st_sub),
        Spacer(1, 3 * mm),
    ]

    hdr = [Paragraph(f"<b>{h}</b>", st_celda) for h in HEADERS]
    data = [hdr]
    for r in filas:
        fila = _fila(r)
        hexcolor = COLORES_ESTADO.get(_estado_de(r), ("111827",))[0]
        fila[5] = f'<font color="#{hexcolor}"><b>{fila[5]}</b></font>'
        data.append([Paragraph(str(v), st_celda) for v in fila])

    tabla = Table(data, colWidths=ANCHOS_PDF, repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{AZUL}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(tabla)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("© Fareas · Facultad de Informática UNAMBA", st_pie))

    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Matriz de asistencia por grupo (R7/11.8): alumnos × fechas con A/T/F/FJ/—
# ---------------------------------------------------------------------------
LEYENDA_MATRIZ = {
    "A": "Asistió",
    "T": "Tardanza",
    "F": "Falta",
    "FJ": "Falta justificada",
    "—": "Sin marca del alumno",
}
COLOR_MATRIZ = {"A": "22C55E", "T": "F59E0B", "F": "EF4444", "FJ": "0EA5E9", "—": "9CA3AF"}


def excel_matriz_grupo(
    filas: list[dict], fechas: list[str], titulo_periodo: str | None = None
) -> bytes:
    """Registro de asistencia del docente (RF-16/11.8): una columna por fecha.

    ``filas``: [{student_name, student_code, celdas: {fecha: 'A'|'T'|'F'|'FJ'|'—'}}].
    Totales y % de asistencia se calculan aquí. Las sesiones canceladas y
    needs_review ya vienen EXCLUIDAS por el router (R4/R5).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz"

    ncols = 2 + len(fechas) + 5
    def _merge(row: int, texto: str, color: str, size: int) -> None:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(ncols, 2))
        c = ws.cell(row=row, column=1, value=texto)
        c.font = Font(bold=True, size=size, color=color)
        c.alignment = Alignment(horizontal="center", vertical="center")

    _merge(1, TITULO, AZUL, 13)
    _merge(2, f"{SUBTITULO} - Registro de asistencia"
           f"{f' ({titulo_periodo})' if titulo_periodo else ''}", "374151", 11)
    _merge(3, f"Generado: {datetime.now():%d/%m/%Y %H:%M}  ·  Alumnos: {len(filas)}  ·  "
              f"Fechas: {len(fechas)}  ·  "
              + "  ".join(f"{k}={v}" for k, v in LEYENDA_MATRIZ.items()), "9CA3AF", 9)

    fila_hdr = 5
    fill_hdr = PatternFill("solid", fgColor=AZUL)
    fino = Side(style="thin", color="D1D5DB")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    headers = ["Estudiante", "Código", *fechas, "A", "T", "F", "FJ", "% Asist."]
    for col, titulo in enumerate(headers, start=1):
        c = ws.cell(row=fila_hdr, column=col, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill_hdr
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = borde

    for i, r in enumerate(filas):
        celda = r.get("celdas", {})
        a = sum(1 for f in fechas if celda.get(f) == "A")
        t = sum(1 for f in fechas if celda.get(f) == "T")
        f_ = sum(1 for f in fechas if celda.get(f) == "F")
        fj = sum(1 for f in fechas if celda.get(f) == "FJ")
        base = a + t + f_ + fj
        pct = round((a + t) * 100 / base, 1) if base else 0.0
        valores = [r.get("student_name", ""), r.get("student_code") or "-",
                   *[celda.get(f, "—") for f in fechas], a, t, f_, fj, pct]
        for col, valor in enumerate(valores, start=1):
            c = ws.cell(row=fila_hdr + 1 + i, column=col, value=valor)
            c.border = borde
            if col > 2 and col <= 2 + len(fechas):
                marca = str(valor)
                c.font = Font(bold=True, color=COLOR_MATRIZ.get(marca, "111827"))
                c.alignment = Alignment(horizontal="center")
            if col == len(headers):
                c.number_format = "0.0"

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 10
    for j in range(len(fechas)):
        ws.column_dimensions[get_column_letter(3 + j)].width = 11
    for j in range(5):
        ws.column_dimensions[get_column_letter(3 + len(fechas) + j)].width = 6
    ws.freeze_panes = f"C{fila_hdr + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Riesgo DPI (R7/11.8): alumnos en alerta/límite por curso
# ---------------------------------------------------------------------------
COLOR_NIVEL = {"alerta": "B45309", "limite": "B91C1C"}


def excel_riesgo_dpi(filas: list[dict], titulo_periodo: str | None = None) -> bytes:
    """Alumnos en riesgo (R7): % de inasistencia, nivel y faltas restantes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Riesgo DPI"

    ws.merge_cells("A1:K1")
    c = ws["A1"]
    c.value = TITULO
    c.font = Font(bold=True, size=13, color=AZUL)
    c.alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:K2")
    c = ws["A2"]
    c.value = f"{SUBTITULO} - Alumnos en riesgo DPI{f' ({titulo_periodo})' if titulo_periodo else ''}"
    c.font = Font(bold=True, size=11, color="374151")
    c.alignment = Alignment(horizontal="center")
    ws.merge_cells("A3:K3")
    ws["A3"] = f"Generado: {datetime.now():%d/%m/%Y %H:%M}  ·  Alumnos: {len(filas)}"
    ws["A3"].font = Font(size=9, color="9CA3AF")
    ws["A3"].alignment = Alignment(horizontal="center")

    headers = ["Estudiante", "Código", "Curso", "Sesiones", "Faltas", "Justificadas",
               "Tardanzas", "% Inasist.", "Nivel", "Permitidas", "Restantes"]
    fila_hdr = 5
    fill_hdr = PatternFill("solid", fgColor=AZUL)
    fino = Side(style="thin", color="D1D5DB")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    for col, titulo in enumerate(headers, start=1):
        c = ws.cell(row=fila_hdr, column=col, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill_hdr
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = borde

    for i, r in enumerate(filas):
        nivel = str(r.get("nivel") or "")
        valores = [
            r.get("student_name", ""), r.get("student_code") or "-",
            f"{r.get('course_code', '')} - {r.get('course_name', '')}",
            r.get("sessions_held"), r.get("absences"), r.get("justified_absences"),
            r.get("lates"), r.get("absence_percent"),
            "LÍMITE (inhabilitado)" if nivel == "limite" else "Alerta",
            r.get("faltas_permitidas") if r.get("faltas_permitidas") is not None else "—",
            r.get("faltas_restantes") if r.get("faltas_restantes") is not None else "—",
        ]
        for col, valor in enumerate(valores, start=1):
            c = ws.cell(row=fila_hdr + 1 + i, column=col, value=valor)
            c.border = borde
            if col == 8:
                c.number_format = "0.00"
            if col == 9:
                c.font = Font(bold=True, color=COLOR_NIVEL.get(nivel, "111827"))

    anchos = [32, 10, 38, 9, 8, 11, 10, 11, 20, 10, 10]
    for col, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.freeze_panes = f"A{fila_hdr + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def nombre_matriz(group_code: str, course_code: str, now: datetime | None = None) -> str:
    return f"asistencia_{course_code}-{group_code}_{(now or datetime.now()):%Y%m%d_%H%M}.xlsx"


def nombre_riesgo(now: datetime | None = None) -> str:
    return f"riesgo_dpi_{(now or datetime.now()):%Y%m%d_%H%M}.xlsx"
