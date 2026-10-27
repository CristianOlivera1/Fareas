import io
import re
import zipfile
import zlib
from dataclasses import dataclass
from datetime import time
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

TITULO = "FACULTAD DE INFORMÁTICA - UNAMBA"
DIAS = {1: "Lunes", 2: "Martes", 3: "Miércoles", 4: "Jueves", 5: "Viernes"}
ROMANOS = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V", 6: "VI", 7: "VII", 8: "VIII", 9: "IX", 10: "X"}
PALETA = ["DBEAFE", "DCFCE7", "FEF3C7", "FCE7F3", "E0E7FF", "CCFBF1", "FFEDD5", "F3E8FF", "ECFCCB", "FEE2E2"]
AZUL = "155DFC"
FORMATOS = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "png": "image/png",
}


@dataclass(frozen=True)
class Celda:
    fila: int
    filas: int
    columna: int
    bloque: dict


@dataclass(frozen=True)
class Rejilla:
    inicio: int
    paso: int
    filas: int
    columnas: tuple[tuple[int, int], ...]
    celdas: tuple[Celda, ...]

    def columnas_del_dia(self, dia: int) -> tuple[int, int]:
        indices = [i for i, (d, _) in enumerate(self.columnas) if d == dia]
        return indices[0], indices[-1]

    def etiqueta(self, fila: int) -> str:
        minutos = self.inicio + fila * self.paso
        return f"{minutos // 60:02d}:{minutos % 60:02d}"


@dataclass(frozen=True)
class Opciones:
    docente: bool = True
    aula: bool = True


def _minutos(valor: time) -> int:
    return valor.hour * 60 + valor.minute


def construir_rejilla(bloques: list[dict]) -> Rejilla:
    if not bloques:
        return Rejilla(480, 60, 10, tuple((d, 0) for d in DIAS), ())
    inicios = [_minutos(b["start_time"]) for b in bloques]
    fines = [_minutos(b["end_time"]) for b in bloques]
    alineado = all(m % 60 == 0 for m in inicios + fines)
    paso = 60 if alineado else 30
    inicio = min(inicios) // paso * paso
    fin = -(-max(fines) // paso) * paso
    filas = max(1, (fin - inicio) // paso)

    carriles: dict[int, list[int]] = {d: [] for d in DIAS}
    asignados: list[tuple[dict, int]] = []
    for bloque in sorted(bloques, key=lambda b: (_minutos(b["start_time"]), b["course_code"])):
        dia = bloque["weekday"]
        if dia not in carriles:
            continue
        ocupados = carriles[dia]
        inicio_b = _minutos(bloque["start_time"])
        for carril, ultimo_fin in enumerate(ocupados):
            if ultimo_fin <= inicio_b:
                ocupados[carril] = _minutos(bloque["end_time"])
                break
        else:
            carril = len(ocupados)
            ocupados.append(_minutos(bloque["end_time"]))
        asignados.append((bloque, carril))

    columnas = tuple((d, c) for d in DIAS for c in range(max(1, len(carriles[d]))))
    indice = {columna: i for i, columna in enumerate(columnas)}
    celdas = tuple(
        Celda(
            fila=(_minutos(b["start_time"]) - inicio) // paso,
            filas=max(1, -(-(_minutos(b["end_time"]) - _minutos(b["start_time"])) // paso)),
            columna=indice[(b["weekday"], carril)],
            bloque=b,
        )
        for b, carril in asignados
    )
    return Rejilla(inicio, paso, filas, columnas, celdas)


def lineas_celda(bloque: dict, opciones: Opciones) -> list[str]:
    lineas = [f"{bloque['course_code']} · {bloque['group_code']}", bloque["course_name"]]
    if opciones.docente and bloque.get("teacher_name"):
        lineas.append(bloque["teacher_name"])
    if opciones.aula and bloque.get("room_code"):
        lineas.append(bloque["room_code"])
    return lineas


def color_curso(codigo: str) -> str:
    return PALETA[zlib.crc32(codigo.encode()) % len(PALETA)]


def romano(ciclo: int) -> str:
    return ROMANOS.get(ciclo, str(ciclo))


def agrupar_por_ciclo(bloques: list[dict]) -> dict[int, list[dict]]:
    ciclos: dict[int, list[dict]] = {}
    for bloque in bloques:
        ciclos.setdefault(bloque["cycle"], []).append(bloque)
    return dict(sorted(ciclos.items()))


def titulo_hoja(semestre: str, ciclo: int) -> str:
    return f"Horario {semestre} · Ciclo {romano(ciclo)}"


def nombre_archivo(semestre: str, ciclo: int | None, extension: str) -> str:
    sufijo = f"ciclo-{romano(ciclo)}" if ciclo is not None else "todos-los-ciclos"
    base = re.sub(r"[^A-Za-z0-9._-]+", "-", f"horario_{semestre}_{sufijo}").strip("-")
    return f"{base}.{extension}"


def _fila_lista(b: dict) -> list:
    return [
        romano(b["cycle"]), b["course_code"], b["course_name"], b["group_code"],
        DIAS.get(b["weekday"], ""), b["start_time"].strftime("%H:%M"), b["end_time"].strftime("%H:%M"),
        b.get("teacher_name") or "", b.get("room_code") or "",
    ]


def excel_horario(bloques: list[dict], semestre: str, opciones: Opciones) -> bytes:
    libro = Workbook()
    libro.remove(libro.active)
    borde = Side(style="thin", color="D1D5DB")
    marco = Border(left=borde, right=borde, top=borde, bottom=borde)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ciclo, filas_ciclo in agrupar_por_ciclo(bloques).items():
        hoja = libro.create_sheet(f"Ciclo {romano(ciclo)}")
        rejilla = construir_rejilla(filas_ciclo)
        total = len(rejilla.columnas) + 1
        hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total)
        hoja.cell(1, 1, f"{TITULO} - {titulo_hoja(semestre, ciclo)}").font = Font(bold=True, size=13)
        hoja.cell(2, 1, "Hora").font = Font(bold=True, color="FFFFFF")
        hoja.cell(2, 1).fill = PatternFill("solid", fgColor=AZUL)
        hoja.cell(2, 1).alignment = centro
        for dia, nombre in DIAS.items():
            primero, ultimo = rejilla.columnas_del_dia(dia)
            if ultimo > primero:
                hoja.merge_cells(start_row=2, start_column=primero + 2, end_row=2, end_column=ultimo + 2)
            celda = hoja.cell(2, primero + 2, nombre)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.alignment = centro
            for columna in range(primero + 2, ultimo + 3):
                hoja.cell(2, columna).fill = PatternFill("solid", fgColor=AZUL)
        for fila in range(rejilla.filas):
            hoja.cell(fila + 3, 1, rejilla.etiqueta(fila)).alignment = centro
            hoja.cell(fila + 3, 1).font = Font(bold=True, color="6B7280")
            hoja.row_dimensions[fila + 3].height = 34 if rejilla.paso == 60 else 28
            for columna in range(2, total + 1):
                hoja.cell(fila + 3, columna).border = marco
        for celda in rejilla.celdas:
            fila, columna = celda.fila + 3, celda.columna + 2
            if celda.filas > 1:
                hoja.merge_cells(
                    start_row=fila, start_column=columna, end_row=fila + celda.filas - 1, end_column=columna
                )
            destino = hoja.cell(fila, columna, "\n".join(lineas_celda(celda.bloque, opciones)))
            destino.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            destino.font = Font(size=9)
            relleno = PatternFill("solid", fgColor=color_curso(celda.bloque["course_code"]))
            for r in range(fila, fila + celda.filas):
                hoja.cell(r, columna).fill = relleno
                hoja.cell(r, columna).border = marco
        hoja.column_dimensions["A"].width = 9
        for columna in range(2, total + 1):
            hoja.column_dimensions[get_column_letter(columna)].width = 26
        hoja.freeze_panes = "B3"
    lista = libro.create_sheet("Lista")
    encabezados = ["Ciclo", "Código", "Curso", "Grupo", "Día", "Inicio", "Fin", "Docente", "Aula"]
    lista.append(encabezados)
    for celda in lista[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
    for bloque in sorted(
        bloques, key=lambda b: (b["cycle"], b["weekday"], b["start_time"], b["course_code"])
    ):
        lista.append(_fila_lista(bloque))
    for indice, ancho in enumerate([7, 12, 34, 8, 12, 9, 9, 30, 12], start=1):
        lista.column_dimensions[get_column_letter(indice)].width = ancho
    lista.freeze_panes = "A2"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def pdf_horario(bloques: list[dict], semestre: str, opciones: Opciones) -> bytes:
    salida = io.BytesIO()
    documento = SimpleDocTemplate(
        salida, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
        topMargin=10 * mm, bottomMargin=10 * mm, title=f"Horario {semestre}",
    )
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("t", parent=estilos["Heading2"], fontSize=12, spaceAfter=1)
    sub = ParagraphStyle("s", parent=estilos["Normal"], fontSize=8, textColor=colors.HexColor("#6B7280"))
    cuerpo = ParagraphStyle("c", parent=estilos["Normal"], fontSize=6.5, leading=7.6, alignment=1)
    hora = ParagraphStyle("h", parent=estilos["Normal"], fontSize=7, alignment=1)
    historia = []
    ciclos = agrupar_por_ciclo(bloques)
    for posicion, (ciclo, filas_ciclo) in enumerate(ciclos.items()):
        rejilla = construir_rejilla(filas_ciclo)
        historia += [
            Paragraph(TITULO, sub),
            Paragraph(titulo_hoja(semestre, ciclo), titulo),
            Spacer(1, 3 * mm),
        ]
        total = len(rejilla.columnas) + 1
        datos = [[""] * total for _ in range(rejilla.filas + 1)]
        datos[0][0] = "Hora"
        estilo = [
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{AZUL}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, 0), 8),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ]
        for dia, nombre in DIAS.items():
            primero, ultimo = rejilla.columnas_del_dia(dia)
            datos[0][primero + 1] = nombre
            if ultimo > primero:
                estilo.append(("SPAN", (primero + 1, 0), (ultimo + 1, 0)))
        for fila in range(rejilla.filas):
            datos[fila + 1][0] = Paragraph(rejilla.etiqueta(fila), hora)
        for celda in rejilla.celdas:
            texto = "<br/>".join(
                f"<b>{escape(linea)}</b>" if i == 0 else escape(linea)
                for i, linea in enumerate(lineas_celda(celda.bloque, opciones))
            )
            columna, fila = celda.columna + 1, celda.fila + 1
            datos[fila][columna] = Paragraph(texto, cuerpo)
            if celda.filas > 1:
                estilo.append(("SPAN", (columna, fila), (columna, fila + celda.filas - 1)))
            estilo.append((
                "BACKGROUND", (columna, fila), (columna, fila + celda.filas - 1),
                colors.HexColor(f"#{color_curso(celda.bloque['course_code'])}"),
            ))
        ancho_util = landscape(A4)[0] - 20 * mm
        ancho_hora = 14 * mm
        ancho_col = (ancho_util - ancho_hora) / len(rejilla.columnas)
        alto_fila = min(34, max(18, 400 / rejilla.filas))
        tabla = Table(
            datos, colWidths=[ancho_hora] + [ancho_col] * len(rejilla.columnas),
            rowHeights=[16] + [alto_fila] * rejilla.filas,
        )
        tabla.setStyle(TableStyle(estilo))
        historia.append(tabla)
        if posicion < len(ciclos) - 1:
            historia.append(PageBreak())
    documento.build(historia)
    return salida.getvalue()


def _fuente(tamano: int, negrita: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    nombres = (
        ["arialbd.ttf", "DejaVuSans-Bold.ttf"] if negrita else ["arial.ttf", "DejaVuSans.ttf"]
    )
    for nombre in nombres:
        try:
            return ImageFont.truetype(nombre, tamano)
        except OSError:
            continue
    return ImageFont.load_default(size=tamano)


def _ajustar(dibujo: ImageDraw.ImageDraw, texto: str, fuente, ancho: int) -> list[str]:
    palabras, lineas, actual = texto.split(), [], ""
    for palabra in palabras:
        prueba = f"{actual} {palabra}".strip()
        if dibujo.textlength(prueba, font=fuente) <= ancho or not actual:
            actual = prueba
        else:
            lineas.append(actual)
            actual = palabra
    if actual:
        lineas.append(actual)
    return lineas


def _imagen_ciclo(filas_ciclo: list[dict], semestre: str, ciclo: int, opciones: Opciones) -> bytes:
    rejilla = construir_rejilla(filas_ciclo)
    ancho_hora, ancho_carril, alto_hora, alto_cab, alto_titulo = 80, 210, 62, 40, 70
    alto_fila = alto_hora if rejilla.paso == 60 else alto_hora // 2 + 8
    ancho = ancho_hora + ancho_carril * len(rejilla.columnas)
    alto = alto_titulo + alto_cab + alto_fila * rejilla.filas
    imagen = Image.new("RGB", (ancho, alto), "white")
    dibujo = ImageDraw.Draw(imagen)
    f_titulo, f_cab, f_hora = _fuente(22, True), _fuente(15, True), _fuente(13)
    f_texto, f_negrita = _fuente(12), _fuente(12, True)
    dibujo.text((16, 14), TITULO, font=_fuente(12), fill="#6B7280")
    dibujo.text((16, 32), titulo_hoja(semestre, ciclo), font=f_titulo, fill="#111827")
    dibujo.rectangle([0, alto_titulo, ancho, alto_titulo + alto_cab], fill=f"#{AZUL}")
    dibujo.text((16, alto_titulo + 11), "Hora", font=f_cab, fill="white")
    for dia, nombre in DIAS.items():
        primero, ultimo = rejilla.columnas_del_dia(dia)
        x0 = ancho_hora + primero * ancho_carril
        x1 = ancho_hora + (ultimo + 1) * ancho_carril
        dibujo.text(
            ((x0 + x1) / 2, alto_titulo + alto_cab / 2), nombre, font=f_cab, fill="white", anchor="mm"
        )
        dibujo.line([x0, alto_titulo, x0, alto], fill="#9CA3AF", width=1)
    y0 = alto_titulo + alto_cab
    for fila in range(rejilla.filas):
        y = y0 + fila * alto_fila
        dibujo.line([0, y, ancho, y], fill="#E5E7EB", width=1)
        dibujo.text((16, y + 6), rejilla.etiqueta(fila), font=f_hora, fill="#6B7280")
    for columna in range(len(rejilla.columnas) + 1):
        x = ancho_hora + columna * ancho_carril
        dibujo.line([x, y0, x, alto], fill="#E5E7EB", width=1)
    for celda in rejilla.celdas:
        x = ancho_hora + celda.columna * ancho_carril + 3
        y = y0 + celda.fila * alto_fila + 3
        x2 = x + ancho_carril - 6
        y2 = y0 + (celda.fila + celda.filas) * alto_fila - 3
        dibujo.rounded_rectangle(
            [x, y, x2, y2], radius=6, fill=f"#{color_curso(celda.bloque['course_code'])}", outline="#9CA3AF"
        )
        cursor = y + 6
        for indice, linea in enumerate(lineas_celda(celda.bloque, opciones)):
            fuente = f_negrita if indice == 0 else f_texto
            for parte in _ajustar(dibujo, linea, fuente, ancho_carril - 22):
                if cursor + 14 > y2:
                    break
                dibujo.text((x + 8, cursor), parte, font=fuente, fill="#111827")
                cursor += 15
    salida = io.BytesIO()
    imagen.save(salida, format="PNG")
    return salida.getvalue()


def png_horario(bloques: list[dict], semestre: str, opciones: Opciones) -> tuple[bytes, str]:
    ciclos = agrupar_por_ciclo(bloques)
    if len(ciclos) == 1:
        (ciclo, filas_ciclo), = ciclos.items()
        return _imagen_ciclo(filas_ciclo, semestre, ciclo, opciones), "image/png"
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as archivo:
        for ciclo, filas_ciclo in ciclos.items():
            archivo.writestr(
                nombre_archivo(semestre, ciclo, "png"), _imagen_ciclo(filas_ciclo, semestre, ciclo, opciones)
            )
    return salida.getvalue(), "application/zip"


def exportar_horario(
    formato: str, bloques: list[dict], semestre: str, ciclo: int | None, opciones: Opciones
) -> tuple[bytes, str, str]:
    if formato == "xlsx":
        return excel_horario(bloques, semestre, opciones), FORMATOS["xlsx"], nombre_archivo(semestre, ciclo, "xlsx")
    if formato == "pdf":
        return pdf_horario(bloques, semestre, opciones), FORMATOS["pdf"], nombre_archivo(semestre, ciclo, "pdf")
    contenido, tipo = png_horario(bloques, semestre, opciones)
    extension = "zip" if tipo == "application/zip" else "png"
    return contenido, tipo, nombre_archivo(semestre, ciclo, extension)
