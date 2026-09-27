"""Dashboard y reportes (Bloque 8): summary, stats, SSE en vivo y RF-16.

Estrategia: el escenario B6 (helpers_b6) crea aula/curso/grupo; aquí se añade
un bloque con el weekday DE HOY para que la sesión quede con session_date =
CURRENT_DATE y los endpoints 'de hoy' tengan datos deterministas.
"""
import json
from datetime import date, datetime, time
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.attendance_engine import PERU_TZ, resolver_sesion
from app.tables import ScheduleBlock
from tests.helpers_b6 import crear, limpiar

API = "/api/v1"


@pytest.fixture
async def escenario(db_session):
    await limpiar(db_session)
    esc = await crear(db_session)
    yield esc
    await limpiar(db_session)


async def _sesion_hoy_con_marcas(db_session, client, admin_headers, esc) -> int:
    """Bloque de HOY + sesión abierta + e1 asistio / e2 tardanza (API RF-21)."""
    hoy = date.today()
    db_session.add(ScheduleBlock(
        group_id=esc["grupo_id"], weekday=hoy.isoweekday(),
        start_time=time(8, 0), end_time=time(10, 0),
    ))
    await db_session.commit()
    ahora = datetime(hoy.year, hoy.month, hoy.day, 8, 15, tzinfo=PERU_TZ)
    sesion, _ = await resolver_sesion(db_session, esc["aula_id"], now=ahora)
    await db_session.commit()
    assert sesion is not None
    for student_id, estado in ((esc["e1"], "asistio"), (esc["e2"], "tardanza")):
        r = await client.post(
            f"{API}/sessions/{sesion.id}/manual-mark",
            headers=admin_headers,
            json={"student_id": student_id, "status": estado},
        )
        assert r.status_code == 201, r.text
    return sesion.id


# ---------------------------------------------------------------------------
# Summary + stats (RF-13, PLAN §5)
# ---------------------------------------------------------------------------
async def test_summary_tarjetas(client, admin_headers, db_session, escenario):
    esc = escenario
    await _sesion_hoy_con_marcas(db_session, client, admin_headers, esc)

    r = await client.get(f"{API}/dashboard/summary", headers=admin_headers)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["asistencias_hoy"] >= 1      # e1
    assert data["tardanzas_hoy"] >= 1        # e2
    assert data["camaras_total"] >= 0
    assert data["alertas_dpi"] >= 0
    assert "at" in data

    # Sin token → 401
    r = await client.get(f"{API}/dashboard/summary")
    assert r.status_code == 401


async def test_stats_hourly_y_distribution(client, admin_headers, db_session, escenario):
    esc = escenario
    await _sesion_hoy_con_marcas(db_session, client, admin_headers, esc)

    r = await client.get(f"{API}/stats/hourly", headers=admin_headers)
    assert r.status_code == 200, r.text
    serie = r.json()
    assert [s["hour"] for s in serie][0] == "07:00"  # franja institucional 07-20
    assert len(serie) == 14
    h08 = next(s for s in serie if s["hour"] == "08:00")
    assert h08["total"] >= 2 and h08["asistio"] >= 1 and h08["tardanza"] >= 1

    # Con day explícito (aunque la sesión sea de hoy)
    r = await client.get(
        f"{API}/stats/hourly", headers=admin_headers, params={"day": date.today().isoformat()}
    )
    assert r.status_code == 200 and r.json()[1]["total"] >= 2  # índice 1 = 08:00

    r = await client.get(
        f"{API}/stats/distribution", headers=admin_headers,
        params={"desde": date.today().isoformat(), "hasta": date.today().isoformat()},
    )
    assert r.status_code == 200, r.text
    dist = r.json()
    assert dist["asistio"] >= 1 and dist["tardanza"] >= 1 and dist["sin_registro"] >= 0


# ---------------------------------------------------------------------------
# SSE en vivo (RF-13)
# ---------------------------------------------------------------------------
async def test_sse_auth(client):
    """SSE: sin token → 422; token basura → 401 (EventSource va por query)."""
    r = await client.get(f"{API}/dashboard/live")
    assert r.status_code == 422  # falta access_token
    r = await client.get(f"{API}/dashboard/live", params={"access_token": "no-es-un-jwt"})
    assert r.status_code == 401


async def test_sse_estudiante_rechazado(client, temp_student):
    from app.core.security import create_access_token

    token = create_access_token(temp_student.id, "estudiante")  # type: ignore[arg-type]
    r = await client.get(f"{API}/dashboard/live", params={"access_token": token})
    assert r.status_code == 403


async def test_sse_envia_veredicto(client, admin_headers):
    """Inyecta un veredicto en camera_loop._ultimo y lo recibe por SSE.

    Consumimos el body_iterator del EventSourceResponse directamente (el
    stream HTTP de httpx no se puede cerrar: la respuesta es infinita).
    """
    import asyncio

    from app.routers.dashboard import dashboard_live
    from app.services import camera_loop

    camera_loop._ultimo[999] = {  # aula falsa solo para el test
        "room_id": 999, "verdict": "asistio", "student_id": 4,
        "similarity": 0.9, "ts": "2026-09-27T08:00:00",
    }
    token = admin_headers["Authorization"].removeprefix("Bearer ")
    try:
        resp = await dashboard_live(access_token=token, room_id=999)
        item = await asyncio.wait_for(resp.body_iterator.__anext__(), timeout=5.0)
        assert item["event"] == "verdict"
        data = json.loads(item["data"])
        assert data["room_id"] == 999 and data["verdict"] == "asistio"
        await resp.body_iterator.aclose()  # corta el generador infinito
    finally:
        camera_loop._ultimo.pop(999, None)


# ---------------------------------------------------------------------------
# Reportes Excel/PDF (RF-16)
# ---------------------------------------------------------------------------
async def test_reporte_excel_y_pdf(client, admin_headers, db_session, escenario):
    esc = escenario
    await _sesion_hoy_con_marcas(db_session, client, admin_headers, esc)
    hoy = date.today().isoformat()

    r = await client.get(
        f"{API}/reports/attendance.xlsx", headers=admin_headers,
        params={"desde": hoy, "hasta": hoy, "course_id": esc["curso_id"]},
    )
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in r.headers.get("content-disposition", "")
    assert r.content[:2] == b"PK"  # zip magic del xlsx
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    assert ws.max_row >= 7  # encabezado (filas 1-5) + 2 registros

    r = await client.get(
        f"{API}/reports/attendance.pdf", headers=admin_headers, params={"desde": hoy, "hasta": hoy}
    )
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert r.content[:5] == b"%PDF-"


async def test_docente_solo_sus_grupos_en_reportes(client, admin_headers, docente_headers,
                                                  db_session, escenario):
    """RF-16 con RF-21: el docente ajeno al grupo obtiene un reporte VACÍO."""
    esc = escenario
    await _sesion_hoy_con_marcas(db_session, client, admin_headers, esc)
    hoy = date.today().isoformat()

    # docente_headers = OTRO docente (temp.docente.login) → 0 filas
    r = await client.get(
        f"{API}/reports/attendance.xlsx", headers=docente_headers,
        params={"desde": hoy, "hasta": hoy, "course_id": esc["curso_id"]},
    )
    assert r.status_code == 200, r.text
    ws = load_workbook(BytesIO(r.content)).active
    assert ws.max_row == 5  # solo encabezado, sin datos

    # admin ve las 2 marcas del curso
    r = await client.get(
        f"{API}/reports/attendance.xlsx", headers=admin_headers,
        params={"desde": hoy, "hasta": hoy, "course_id": esc["curso_id"]},
    )
    ws = load_workbook(BytesIO(r.content)).active
    assert ws.max_row >= 7
