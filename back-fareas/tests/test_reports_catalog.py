import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app.services.report_render import (
    Celda,
    Columna,
    Indicador,
    Tabla,
    a_csv,
    a_pdf,
    a_xlsx,
    nombre_archivo,
)
from tests.test_dashboard_reports import _sesion_hoy_con_marcas, escenario  # noqa: F401

API = "/api/v1"


def _tabla(filas: int = 3) -> Tabla:
    return Tabla(
        "asistencia-detallada", "Asistencia detallada", "Descripción del reporte",
        [Columna("a", "Estudiante", 2), Columna("b", "Estado", 1, "center"), Columna("c", "Pct", 1, "right", True)],
        [[Celda(f"Alumno {i}"), Celda("Falta", "bad"), Celda(80.5)] for i in range(filas)],
        [("Periodo", "01/01/2026 al 31/01/2026")],
        [Indicador("Registros", str(filas), "muted"), Indicador("% de asistencia", "80%", "ok")],
    )


def test_csv_has_bom_header_and_rows():
    contenido = a_csv(_tabla()).decode("utf-8-sig").splitlines()

    assert contenido[0] == "Estudiante,Estado,Pct"
    assert contenido[1] == "Alumno 0,Falta,80.5"
    assert len(contenido) == 4


def test_xlsx_has_institutional_header_filters_and_table():
    hoja = load_workbook(io.BytesIO(a_xlsx(_tabla()))).active

    assert "UNAMBA" in hoja["A1"].value
    assert hoja["A2"].value == "Asistencia detallada"
    assert "Periodo" in hoja["A4"].value
    textos = [c.value for fila in hoja.iter_rows() for c in fila]
    assert "Estudiante" in textos and "Alumno 2" in textos


def test_pdf_renders_even_without_rows_and_with_many_pages():
    assert a_pdf(_tabla(0)).startswith(b"%PDF")
    grande = a_pdf(_tabla(300))
    assert grande.startswith(b"%PDF") and grande.count(b"/Type /Page\n") + grande.count(b"/Type /Page ") > 2


def test_file_names_are_slugged_and_typed():
    assert nombre_archivo(_tabla(), "csv").startswith("asistencia-detallada_")
    assert nombre_archivo(_tabla(), "pdf").endswith(".pdf")


@pytest.mark.asyncio
async def test_catalog_lists_the_reports(client, admin_headers):
    r = await client.get(f"{API}/reports/catalog", headers=admin_headers)

    assert r.status_code == 200
    assert {item["key"] for item in r.json()} == {
        "attendance", "course-summary", "student-summary", "dpi-risk",
    }
    assert (await client.get(f"{API}/reports/catalog")).status_code == 401


@pytest.mark.asyncio
async def test_preview_and_export_with_filters(client, admin_headers, db_session, escenario):  # noqa: F811
    esc = escenario
    await _sesion_hoy_con_marcas(db_session, client, admin_headers, esc)
    hoy = date.today().isoformat()
    base = {"desde": hoy, "hasta": hoy, "course_id": esc["curso_id"]}

    detalle = (await client.get(
        f"{API}/reports/preview", headers=admin_headers, params={**base, "kind": "attendance"},
    )).json()
    assert detalle["total"] >= 2
    assert {i["label"] for i in detalle["indicators"]} >= {"Registros", "% de asistencia", "Faltas"}
    assert any(f["label"] == "Curso" for f in detalle["filters"])

    solo_tardanza = (await client.get(
        f"{API}/reports/preview", headers=admin_headers,
        params={**base, "kind": "attendance", "status": "tardanza"},
    )).json()
    assert 1 <= solo_tardanza["total"] < detalle["total"]
    assert all(fila[8]["value"] == "Tardanza" for fila in solo_tardanza["rows"])

    for kind in ("course-summary", "student-summary"):
        for formato, firma in (("pdf", b"%PDF"), ("xlsx", b"PK"), ("csv", b"\xef\xbb\xbf")):
            r = await client.get(
                f"{API}/reports/export", headers=admin_headers,
                params={**base, "kind": kind, "format": formato},
            )
            assert r.status_code == 200, r.text
            assert r.content.startswith(firma)
            assert "attachment" in r.headers["content-disposition"]

    riesgo = await client.get(
        f"{API}/reports/export", headers=admin_headers,
        params={"kind": "dpi-risk", "format": "xlsx", "course_id": esc["curso_id"]},
    )
    assert riesgo.status_code == 200


@pytest.mark.asyncio
async def test_teacher_only_sees_own_groups_and_bad_input_is_rejected(
    client, admin_headers, docente_headers, db_session, escenario,  # noqa: F811
):
    esc = escenario
    await _sesion_hoy_con_marcas(db_session, client, admin_headers, esc)
    hoy = date.today().isoformat()

    propio = await client.get(
        f"{API}/reports/preview", headers=docente_headers,
        params={"kind": "attendance", "desde": hoy, "hasta": hoy, "course_id": esc["curso_id"]},
    )
    assert propio.status_code == 200
    assert propio.json()["total"] == 2

    invalidos = (
        {"kind": "attendance", "desde": "2026-02-01", "hasta": "2026-01-01"},
        {"kind": "attendance", "status": "inventado"},
    )
    for params in invalidos:
        assert (await client.get(f"{API}/reports/preview", headers=admin_headers, params=params)).status_code == 422
    assert (await client.get(f"{API}/reports/preview", headers=admin_headers, params={"kind": "nope"})).status_code == 404


@pytest.mark.asyncio
async def test_every_report_runs_with_every_filter_on_the_real_schema(client, admin_headers, docente_headers):
    filtros = {
        "desde": "2020-01-01", "hasta": date.today().isoformat(), "cycle": 9, "course_id": 1,
        "group_id": 1, "teacher_id": 1, "room_id": 1, "status": "asistio,tardanza,falta,falta_justificada",
        "q": "a", "only_risk": "true",
    }
    for headers in (admin_headers, docente_headers):
        for kind in ("attendance", "course-summary", "student-summary", "dpi-risk"):
            for params in ({"kind": kind}, {"kind": kind, **filtros}):
                r = await client.get(f"{API}/reports/preview", headers=headers, params=params)
                assert r.status_code == 200, (kind, r.text)
                assert r.json()["columns"]
            for formato in ("pdf", "xlsx", "csv"):
                r = await client.get(
                    f"{API}/reports/export", headers=headers, params={"kind": kind, "format": formato},
                )
                assert r.status_code == 200, (kind, formato, r.text)
