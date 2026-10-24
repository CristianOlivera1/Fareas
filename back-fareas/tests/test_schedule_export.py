import io
import zipfile
from datetime import time

import pytest
from openpyxl import load_workbook
from sqlmodel import delete, select

from app.services.schedule_export import (
    Opciones,
    construir_rejilla,
    exportar_horario,
    lineas_celda,
    nombre_archivo,
)
from app.tables import Course, CourseGroup, ScheduleBlock

API = "/api/v1"
CURSO = "TSTEX1"


def _bloque(codigo="ISA901", dia=1, inicio=(7, 0), fin=(9, 0), grupo="A", ciclo=9, **extra) -> dict:
    return {
        "course_code": codigo, "course_name": "Computación Gráfica", "group_code": grupo,
        "teacher_name": "MAMANI VILCA, Eddy", "room_code": "LAB 401", "cycle": ciclo,
        "weekday": dia, "start_time": time(*inicio), "end_time": time(*fin), **extra,
    }


def test_the_grid_spans_two_hour_blocks_over_two_rows():
    rejilla = construir_rejilla([_bloque(), _bloque("ISA902", dia=2, inicio=(8, 0), fin=(9, 0))])

    assert (rejilla.inicio, rejilla.paso, rejilla.filas) == (420, 60, 2)
    assert [(c.fila, c.filas) for c in rejilla.celdas] == [(0, 2), (1, 1)]


def test_overlapping_groups_get_their_own_column_in_the_same_day():
    rejilla = construir_rejilla([_bloque(grupo="A"), _bloque(grupo="B")])

    lunes = rejilla.columnas_del_dia(1)
    assert lunes[1] - lunes[0] == 1
    assert len({c.columna for c in rejilla.celdas}) == 2


def test_half_hour_blocks_switch_to_half_hour_rows():
    rejilla = construir_rejilla([_bloque(inicio=(7, 30), fin=(9, 0))])

    assert rejilla.paso == 30 and rejilla.filas == 3


def test_cell_lines_respect_the_teacher_and_room_options():
    assert len(lineas_celda(_bloque(), Opciones())) == 4
    assert lineas_celda(_bloque(), Opciones(docente=False, aula=False)) == [
        "ISA901 · A", "Computación Gráfica",
    ]


def test_file_names_are_safe():
    assert nombre_archivo("2026-II", 4, "pdf") == "horario_2026-II_ciclo-IV.pdf"
    assert nombre_archivo("2026 II/x", None, "xlsx") == "horario_2026-II-x_todos-los-ciclos.xlsx"


def test_excel_has_one_sheet_per_cycle_and_a_list():
    bloques = [_bloque(), _bloque("ISA801", ciclo=8, dia=3)]
    contenido, tipo, nombre = exportar_horario("xlsx", bloques, "2026-II", None, Opciones())

    libro = load_workbook(io.BytesIO(contenido))
    assert libro.sheetnames == ["Ciclo VIII", "Ciclo IX", "Lista"]
    assert nombre.endswith("todos-los-ciclos.xlsx") and "spreadsheetml" in tipo
    assert libro["Lista"].max_row == 3


def test_pdf_and_png_are_valid_files():
    pdf, tipo_pdf, _ = exportar_horario("pdf", [_bloque()], "2026-II", 9, Opciones())
    png, tipo_png, nombre = exportar_horario("png", [_bloque()], "2026-II", 9, Opciones())

    assert pdf.startswith(b"%PDF") and tipo_pdf == "application/pdf"
    assert png.startswith(b"\x89PNG") and tipo_png == "image/png" and nombre.endswith(".png")


def test_png_of_several_cycles_is_a_zip():
    contenido, tipo, nombre = exportar_horario(
        "png", [_bloque(), _bloque("ISA801", ciclo=8)], "2026-II", None, Opciones()
    )

    assert tipo == "application/zip" and nombre.endswith(".zip")
    assert len(zipfile.ZipFile(io.BytesIO(contenido)).namelist()) == 2


@pytest.mark.asyncio
async def test_endpoint_exports_the_filtered_schedule(client, admin_headers, db_session):
    antiguos = (await db_session.exec(select(Course).where(Course.code == CURSO))).all()
    for curso in antiguos:
        ids = (await db_session.exec(select(CourseGroup.id).where(CourseGroup.course_id == curso.id))).all()
        if ids:
            await db_session.exec(delete(ScheduleBlock).where(ScheduleBlock.group_id.in_(ids)))
            await db_session.exec(delete(CourseGroup).where(CourseGroup.id.in_(ids)))
        await db_session.exec(delete(Course).where(Course.id == curso.id))
    await db_session.commit()
    try:
        curso = (await client.post(
            f"{API}/courses", headers=admin_headers,
            json={"code": CURSO, "name": "Curso de exportación", "cycle": 9},
        )).json()
        grupo = (await client.post(
            f"{API}/courses/{curso['id']}/groups", headers=admin_headers, json={"group_code": "A"},
        )).json()
        creado = await client.post(
            f"{API}/groups/{grupo['id']}/blocks", headers=admin_headers,
            json={"weekday": 1, "start_time": "07:00:00", "end_time": "09:00:00"},
        )
        assert creado.status_code == 201, creado.text

        for formato, firma in (("pdf", b"%PDF"), ("xlsx", b"PK"), ("png", b"\x89PNG")):
            r = await client.get(
                f"{API}/schedule/export", headers=admin_headers,
                params={"format": formato, "cycle": 9, "course_id": curso["id"]},
            )
            assert r.status_code == 200, r.text
            assert r.content.startswith(firma)
            assert "attachment" in r.headers["content-disposition"]

        vacio = await client.get(
            f"{API}/schedule/export", headers=admin_headers,
            params={"format": "pdf", "semester": "1900-I"},
        )
        assert vacio.status_code == 404
    finally:
        ids = (await db_session.exec(select(Course.id).where(Course.code == CURSO))).all()
        for course_id in ids:
            gids = (await db_session.exec(select(CourseGroup.id).where(CourseGroup.course_id == course_id))).all()
            if gids:
                await db_session.exec(delete(ScheduleBlock).where(ScheduleBlock.group_id.in_(gids)))
                await db_session.exec(delete(CourseGroup).where(CourseGroup.id.in_(gids)))
            await db_session.exec(delete(Course).where(Course.id == course_id))
        await db_session.commit()
